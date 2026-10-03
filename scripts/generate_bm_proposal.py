"""Generate the Watt business model proposal as a DOCX document."""

from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "Watt_BM_요금제_제안서_2026-10-04.docx"
NAVY = RGBColor(26, 45, 69)
BLUE = RGBColor(38, 99, 167)
MUTED = RGBColor(88, 103, 119)


def set_font(run, size=10, bold=False, color=NAVY):
    run.font.name = "맑은 고딕"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "맑은 고딕")
    return run


def add_text(paragraph, value, size=10, bold=False, color=NAVY):
    return set_font(paragraph.add_run(value), size, bold, color)


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def heading(doc, title):
    p = doc.add_paragraph(style="Heading 1")
    add_text(p, title, 13, True, BLUE)
    return p


def body(doc, text=""):
    p = doc.add_paragraph()
    add_text(p, text)
    return p


def bullet(doc, label, detail):
    p = doc.add_paragraph(style="List Bullet")
    add_text(p, label, 10, True)
    add_text(p, detail)
    return p


def table(doc, headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.style = "Table Grid"
    for i, title in enumerate(headers):
        cell = t.rows[0].cells[i]
        shade(cell, "DDEAF6")
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        add_text(cell.paragraphs[0], title, 9, True, NAVY)
    for row_index, row in enumerate(rows):
        cells = t.add_row().cells
        for i, value in enumerate(row):
            if row_index % 2 == 1:
                shade(cells[i], "F6F9FC")
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            add_text(cells[i].paragraphs[0], str(value), 8.5)
    if widths:
        for row in t.rows:
            for i, width in enumerate(widths):
                row.cells[i].width = Cm(width)
    doc.add_paragraph()
    return t


def link(paragraph, text, url):
    part = paragraph.part
    relationship_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "2663A7")
    properties.append(color)
    run.append(properties)
    text_element = OxmlElement("w:t")
    text_element.text = text
    run.append(text_element)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


doc = Document()
section = doc.sections[0]
section.top_margin = Cm(1.7)
section.bottom_margin = Cm(1.6)
section.left_margin = Cm(1.8)
section.right_margin = Cm(1.8)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "맑은 고딕"
normal.font.size = Pt(10)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.18
for style_name in ("Heading 1", "Heading 2"):
    styles[style_name].paragraph_format.space_before = Pt(11)
    styles[style_name].paragraph_format.space_after = Pt(4)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
add_text(p, "Watt 요금제 및 대학 제휴 제안서", 18, True)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
add_text(p, "무료 · 개인 구독 · 팀 구독 · 대학 제휴  |  2026년 10월 4일", 9, False, MUTED)

body(doc, "이 문서는 현재 앱의 팀 과제 흐름을 바탕으로 한 초기 가격 실험안입니다. 아래 금액은 공개 판매가나 확정 견적이 아니며, 모두 부가세 포함 가정입니다.")

heading(doc, "1. 권장 요금제 한눈에 보기")
table(
    doc,
    ["요금제", "제안 가격", "이용 단위", "AI 제공량", "핵심 대상"],
    [
        ["무료", "0원", "개인 계정", "가입 시 8크레딧 1회", "처음 과제를 시작하는 학생"],
        ["개인 구독", "월 5,900원\n4개월 19,900원", "개인 계정", "월 30크레딧", "혼자 기획·발표를 준비하는 학생"],
        ["팀 구독", "월 14,900원\n4개월 49,900원", "방 1개 / 최대 5명", "방 공동 월 100크레딧", "한 팀이 비용을 나눠 쓰는 과제"],
        ["대학 제휴", "1인·학기 8,000원\n최소 200명: 160만원", "수업·학과 계약", "1인·학기 16크레딧", "학교가 비용을 부담하는 수업"],
    ],
    [2.6, 3.1, 3.2, 3.2, 4.5],
)
body(doc, "월 구독은 매월 갱신됩니다. 4개월 상품은 한 학기용 기간제 이용권으로 제안하며 자동 갱신하지 않습니다. AI 제공량은 비용 측정 후 조정할 실험값입니다.")

heading(doc, "2. 무엇을 무료로 두고, 어디에 과금할까")
bullet(doc, "무료 핵심: ", "아이디어 직접 입력, 방 초대, 마인드맵, 팀 평가, 진행 상태 확인, 기존 결과 열람과 기본 내보내기. 협업 자체를 막지 않아 팀 전체가 앱을 사용할 수 있게 합니다.")
bullet(doc, "유료 핵심: ", "텍스트·이미지에서 AI 아이디어 추출, AI 비교·분석, MVP 계획 생성, 발표자료 생성·재생성의 추가 이용량. 실제 서버 비용과 시간 절감 가치가 생기는 지점입니다.")
bullet(doc, "내보내기 원칙: ", "사용자가 작성한 내용과 이미 생성된 결과의 기본 DOCX·PDF·PPTX 저장은 구독 만료 뒤에도 허용합니다. 유료 전용 템플릿은 별도 기능을 만든 뒤 검토합니다.")

heading(doc, "3. AI 크레딧 사용 규칙 — 제안")
table(
    doc,
    ["AI 작업", "차감"],
    [
        ["텍스트 아이디어 추출 또는 아이디어 비교·초안 진단", "1크레딧 / 실행"],
        ["이미지 글자 추출·아이디어 추출 또는 MVP 계획 생성", "2크레딧 / 실행"],
        ["발표자료 최초 생성 또는 전체 재생성", "4크레딧 / 실행"],
    ],
    [12.4, 4.2],
)
body(doc, "예시: 텍스트 추출 1회 + 비교 1회 + MVP 1회 + 발표자료 1회 = 8크레딧. 무료 체험으로 과제 흐름을 한 번 완주할 수 있습니다. 실패한 서버 요청은 차감하지 않고, 월 크레딧은 이월하지 않는 안을 권합니다.")

heading(doc, "4. 요금제별 제공 범위")
table(
    doc,
    ["항목", "무료", "개인 구독", "팀 구독", "대학 제휴"],
    [
        ["활성 과제", "1개", "3개", "방당 10개", "계약 수업 범위"],
        ["협업", "최대 5명 방 참여", "무료와 동일", "최대 5명 공동 이용", "수업 참여자 공동 이용"],
        ["AI 이용량", "가입 시 8크레딧", "월 30크레딧", "방 공동 월 100크레딧", "1인·학기 16크레딧"],
        ["결제 책임", "없음", "구독자", "방장 등 결제자 1명", "대학·학과·교수 측"],
        ["학교 관리", "없음", "없음", "없음", "수업별 이용 현황·초대 관리(개발 필요)"],
    ],
    [2.8, 3.2, 3.2, 3.5, 3.9],
)
body(doc, "개인 구독의 AI 크레딧은 다른 팀원에게 이전하지 않습니다. 팀 구독은 방에 귀속되므로 결제자 한 명만 유료여도 방의 모든 멤버가 공동 한도를 사용합니다. 기존 과제의 읽기·내보내기는 해지 뒤에도 유지하는 설계를 권합니다.")

heading(doc, "5. 대학생 인증 할인과 대학 제휴")
table(
    doc,
    ["혜택", "일반 가격", "대학생 인증 가격"],
    [
        ["개인 월 구독", "5,900원", "3,900원"],
        ["개인 4개월 이용권", "19,900원", "13,900원"],
        ["팀 월 구독", "14,900원", "11,900원"],
        ["팀 4개월 이용권", "49,900원", "39,900원"],
    ],
    [6.1, 5.2, 5.3],
)
body(doc, "학생 인증은 학교 이메일 인증과 1년 주기 재확인으로 시작합니다. 팀 할인은 결제자에게 학생 인증이 있을 때 적용합니다. 학교 계약에 포함된 학생은 계약 범위의 수업에서 별도 결제 없이 사용하며, 개인 구독과 학교 제공량은 중복 적립하지 않는 안을 권합니다.")
body(doc, "대학 제휴 기준가는 학생 1인당 4개월 학기 8,000원, 최소 200명 계약(160만원)입니다. 예를 들어 300명 수업은 학기 240만원입니다. 50~100명 규모의 유료 파일럿은 별도 정액 견적으로 운영하고, 정식 확대 전에 학교가 요구하는 관리·보안 기능을 확인합니다.")

heading(doc, "6. 가격 검증과 출시 순서")
bullet(doc, "1단계 — 원가 확인: ", "기능별 AI 호출비, 재시도율, 평균 월 사용 횟수를 기록합니다. 목표는 결제 수수료와 AI·서버 비용을 뺀 매출총이익률 60% 이상입니다. 한도가 원가에 맞지 않으면 가격보다 크레딧 구성을 먼저 조정합니다.")
bullet(doc, "2단계 — 소규모 가격 실험: ", "개인 월 5,900원과 팀 월 14,900원을 시작 가격으로 20~30개 실제 팀에 제시합니다. 무료→유료 전환, 유료 팀의 크레딧 소진, 과제 완료율, 해지 이유를 봅니다.")
bullet(doc, "3단계 — 학교 파일럿: ", "1~2개 수업에서 교수·조교의 초대 및 이용 현황 요구를 확인합니다. 익명 집계 중심으로 설계하고, 학생 과제 내용의 개별 열람 권한은 별도 합의가 있을 때만 제공합니다.")

heading(doc, "7. 현재 구현 상태와 출시 전 작업")
body(doc, "현재 앱에는 과제 단계, 방 협업, AI Edge Functions, 발표자료 재생성·버전 비교, DOCX·PDF·PPTX 내보내기가 있습니다. 결제, 구독 권한, 사용량 원장, 학생 인증, 학교 관리자 화면은 이 제안서의 신규 개발 범위입니다. AI 한도는 화면 표시만으로 처리하지 말고 Edge Function 서버에서 방·사용자 권한과 사용량을 확인해야 합니다.")
body(doc, "모바일 앱에서 개인·팀 디지털 구독을 판매할 때는 Apple과 Google Play의 결제 정책을 채널별로 검토해야 합니다. 대학 기관 계약의 이용 권한 제공은 소비자 판매와 조건이 다를 수 있으므로 계약·배포 구조를 정한 뒤 확인합니다.")

heading(doc, "참고 자료")
sources = [
    ("Expo SDK 56 문서", "https://docs.expo.dev/versions/v56.0.0/"),
    ("Notion 교육 요금제 — 학교 이메일 및 재인증 사례", "https://www.notion.com/help/notion-for-education"),
    ("Figma for Education — 교육용 무료 제공과 AI 크레딧 사례", "https://help.figma.com/hc/en-us/articles/360041061214-Figma-for-Education"),
    ("Canva Campus — 대학 기관 계약 사례", "https://www.canva.com/en_gb/higher-education/faculty/"),
    ("Apple App Review Guidelines", "https://developer.apple.com/app-store/review/guidelines/"),
    ("Google Play Payments policy", "https://support.google.com/googleplay/android-developer/answer/9858738?hl=en"),
]
for name, url in sources:
    p = doc.add_paragraph(style="List Bullet")
    link(p, name, url)

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
add_text(footer, "Watt | BM 가격 실험안 · 2026.10.04", 8, False, MUTED)

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUTPUT)
print(OUTPUT)
